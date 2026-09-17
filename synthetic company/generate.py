#!/usr/bin/env python3
"""Render the Aldermoor Technologies corpus into PDF, DOCX and Markdown.

The documents in this folder are generated, not hand-maintained: the source
of truth is ``content.py``. Regenerate after editing it with::

    uv run --with reportlab --with python-docx python "generate.py"

Only three formats are produced, because those are the only three the
ingestion pipeline accepts — see ``_UPLOAD_MIME_TO_FILE_TYPE`` in
``backend/src/ai_customer_assistant/api/ingest.py``. A .txt or .xlsx file
here would be rejected at upload with HTTP 415, so generating one would
only create something that looks ingestible and is not.

Block types, shared by all three renderers:

    ("h1"|"h2"|"h3", text)   a heading
    ("p", text)              a paragraph
    ("bullets", [text, ...]) a bulleted list
    ("table", [[cell, ...]]) a table whose first row is the header
"""

from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from content import COMPANY, DOCUMENTS, TAGLINE, WEBSITE  # noqa: E402

OUT = {"pdf": HERE / "pdf", "docx": HERE / "docx", "md": HERE / "markdown"}


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def render_pdf(doc: dict, path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    ink = colors.HexColor("#1a1f2b")
    accent = colors.HexColor("#3b5bdb")
    rule = colors.HexColor("#d5dae3")
    band = colors.HexColor("#f2f4f8")

    base = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle("t", parent=base["Title"], fontSize=23, leading=28,
                                textColor=ink, alignment=TA_LEFT, spaceAfter=4),
        "subtitle": ParagraphStyle("st", parent=base["Normal"], fontSize=11,
                                   leading=15, textColor=colors.HexColor("#5b6478"),
                                   spaceAfter=18),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontSize=15, leading=19,
                             textColor=accent, spaceBefore=16, spaceAfter=7),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontSize=12.5, leading=16,
                             textColor=ink, spaceBefore=13, spaceAfter=5),
        "h3": ParagraphStyle("h3", parent=base["Heading3"], fontSize=11, leading=14,
                             textColor=ink, spaceBefore=10, spaceAfter=3),
        "p": ParagraphStyle("p", parent=base["BodyText"], fontSize=9.8, leading=14.5,
                            textColor=ink, spaceAfter=7),
        "li": ParagraphStyle("li", parent=base["BodyText"], fontSize=9.8, leading=14,
                             textColor=ink, spaceAfter=2),
        "cell": ParagraphStyle("c", parent=base["BodyText"], fontSize=8.6, leading=11.6,
                               textColor=ink, spaceAfter=0),
        "cellhead": ParagraphStyle("ch", parent=base["BodyText"], fontSize=8.6,
                                   leading=11.6, textColor=colors.white,
                                   fontName="Helvetica-Bold", spaceAfter=0),
        "foot": ParagraphStyle("f", parent=base["Normal"], fontSize=7.6,
                               textColor=colors.HexColor("#8a93a6")),
    }

    def footer(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(rule)
        canvas.setLineWidth(0.5)
        canvas.line(20 * mm, 15 * mm, A4[0] - 20 * mm, 15 * mm)
        canvas.setFont("Helvetica", 7.6)
        canvas.setFillColor(colors.HexColor("#8a93a6"))
        canvas.drawString(20 * mm, 10.5 * mm, f"{COMPANY} — {doc['title']}")
        canvas.drawRightString(A4[0] - 20 * mm, 10.5 * mm, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    story = [Paragraph(doc["title"], styles["title"])]
    if doc.get("subtitle"):
        story.append(Paragraph(doc["subtitle"], styles["subtitle"]))

    for kind, payload in doc["blocks"]:
        if kind in ("h1", "h2", "h3"):
            story.append(Paragraph(payload, styles[kind]))
        elif kind == "p":
            story.append(Paragraph(payload, styles["p"]))
        elif kind == "bullets":
            story.append(ListFlowable(
                [ListItem(Paragraph(t, styles["li"]), leftIndent=14)
                 for t in payload],
                bulletType="bullet", start="•",
                bulletFontName="Helvetica", bulletFontSize=9, bulletOffsetY=-1,
                leftIndent=14, spaceAfter=9,
            ))
        elif kind == "table":
            header, *rows = payload
            data = [[Paragraph(str(c), styles["cellhead"]) for c in header]]
            data += [[Paragraph(str(c), styles["cell"]) for c in r] for r in rows]
            avail = A4[0] - 40 * mm
            table = Table(data, colWidths=[avail / len(header)] * len(header),
                          repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), accent),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, band]),
                ("GRID", (0, 0), (-1, -1), 0.4, rule),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.extend([table, Spacer(1, 9)])
        elif kind == "pagebreak":
            story.append(PageBreak())

    story.append(Spacer(1, 14))
    story.append(Paragraph(
        f"{COMPANY} — {TAGLINE} — {WEBSITE}. This document describes a company "
        "that does not exist and is provided as sample data.", styles["foot"]))

    SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=18 * mm, bottomMargin=22 * mm,
        title=doc["title"], author=COMPANY, subject=doc.get("subtitle", ""),
    ).build(story, onFirstPage=footer, onLaterPages=footer)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def render_docx(doc: dict, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    d = Document()
    for name, size in (("Normal", 10.5), ("Heading 1", 14), ("Heading 2", 12)):
        try:
            d.styles[name].font.name = "Calibri"
            d.styles[name].font.size = Pt(size)
        except KeyError:
            pass

    d.core_properties.title = doc["title"]
    d.core_properties.author = COMPANY
    d.core_properties.comments = "Synthetic sample data — this company does not exist."

    d.add_heading(doc["title"], level=0)
    if doc.get("subtitle"):
        run = d.add_paragraph().add_run(doc["subtitle"])
        run.italic = True
        run.font.color.rgb = RGBColor(0x5B, 0x64, 0x78)

    for kind, payload in doc["blocks"]:
        if kind == "h1":
            d.add_heading(payload, level=1)
        elif kind == "h2":
            d.add_heading(payload, level=2)
        elif kind == "h3":
            d.add_heading(payload, level=3)
        elif kind == "p":
            d.add_paragraph(payload)
        elif kind == "bullets":
            for item in payload:
                d.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            header, *rows = payload
            table = d.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, header):
                cell.text = str(text)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
            for row in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = str(text)
            d.add_paragraph()

    closing = d.add_paragraph().add_run(
        f"{COMPANY} — {WEBSITE}. This document describes a company that does "
        "not exist and is provided as sample data.")
    closing.italic = True
    closing.font.size = Pt(8)
    closing.font.color.rgb = RGBColor(0x8A, 0x93, 0xA6)

    d.save(str(path))


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def render_md(doc: dict, path: Path) -> None:
    out = [f"# {doc['title']}", ""]
    if doc.get("subtitle"):
        out += [f"*{doc['subtitle']}*", ""]

    for kind, payload in doc["blocks"]:
        if kind == "h1":
            out += [f"## {payload}", ""]
        elif kind == "h2":
            out += [f"### {payload}", ""]
        elif kind == "h3":
            out += [f"#### {payload}", ""]
        elif kind == "p":
            out += [payload, ""]
        elif kind == "bullets":
            out += [f"- {item}" for item in payload] + [""]
        elif kind == "table":
            header, *rows = payload
            out += ["| " + " | ".join(str(c) for c in header) + " |",
                    "|" + "|".join("---" for _ in header) + "|"]
            out += ["| " + " | ".join(str(c).replace("|", "\\|") for c in r) + " |"
                    for r in rows]
            out += [""]

    out += ["---", "",
            f"*{COMPANY} — {TAGLINE} — {WEBSITE}. This document describes a "
            "company that does not exist and is provided as sample data.*", ""]
    path.write_text("\n".join(out), encoding="utf-8")


# ---------------------------------------------------------------------------

RENDERERS = {"pdf": render_pdf, "docx": render_docx, "md": render_md}
EXTENSION = {"pdf": ".pdf", "docx": ".docx", "md": ".md"}


def main() -> int:
    for directory in OUT.values():
        directory.mkdir(parents=True, exist_ok=True)

    for doc in DOCUMENTS:
        fmt = doc["format"]
        path = OUT[fmt] / (doc["slug"] + EXTENSION[fmt])
        RENDERERS[fmt](doc, path)
        print(f"  {path.relative_to(HERE)}  ({path.stat().st_size:,} bytes)")

    print(f"\n{len(DOCUMENTS)} documents generated for {COMPANY}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
